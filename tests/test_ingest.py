"""Conversion to markdown, including a second run that skips fresh files."""

from pathlib import Path

import pytest

from beyond_vector_search.ingest import html_to_markdown, ingest

ROOT = Path(__file__).resolve().parents[1]


def test_html_keeps_headings_and_list_items(tmp_path: Path) -> None:
    source = tmp_path / "raw"
    source.mkdir()
    (source / "notice.html").write_text(
        (ROOT / "examples" / "raw" / "notice.html").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    dest = tmp_path / "docs"
    first = ingest(source, dest)
    assert first[0].status == "wrote"
    text = (dest / "notice.md").read_text(encoding="utf-8")
    assert "# Visitor wifi" in text
    assert "Monday" in text
    assert "- Do not connect lab hardware" in text

    second = ingest(source, dest)
    assert second[0].status == "skipped"


def test_destination_inside_the_source_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "raw"
    source.mkdir()
    with pytest.raises(ValueError):
        ingest(source, source / "out")


def test_html_helper_on_a_fragment() -> None:
    rendered = html_to_markdown(ROOT / "examples" / "raw" / "notice.html")
    assert "guest network" in rendered


def test_blank_pdf_is_marked_for_ocr(tmp_path: Path) -> None:
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    from beyond_vector_search.ingest import pdf_to_markdown

    pdf = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)

    text = pdf_to_markdown(pdf)
    assert "# scan" in text
    assert "## Page 1" in text
    assert "needs OCR" in text


def test_docx_headings_survive(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from beyond_vector_search.ingest import docx_to_markdown

    path = tmp_path / "addendum.docx"
    document = Document()
    document.add_heading("Liability cap", level=2)
    document.add_paragraph("The cap is 2x annual fees, except as set out in MSA §9.3.")
    document.save(path)

    text = docx_to_markdown(path)
    assert "## Liability cap" in text
    assert "2x annual fees" in text


def test_docx_table_becomes_a_markdown_table(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from beyond_vector_search.ingest import docx_to_markdown

    path = tmp_path / "fees.docx"
    document = Document()
    document.add_paragraph("Fees")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Amount"
    table.rows[1].cells[0].text = "Cap"
    table.rows[1].cells[1].text = "2x annual fees"
    document.save(path)

    text = docx_to_markdown(path)
    assert "| Item | Amount |" in text
    assert "| --- | --- |" in text
    assert "| Cap | 2x annual fees |" in text


def test_layout_text_keeps_a_table_and_leaves_prose_alone() -> None:
    from beyond_vector_search.ingest import layout_to_markdown

    table = layout_to_markdown("Item     Amount\nCap      2x annual fees\n")
    assert "| Item | Amount |" in table
    assert "| Cap | 2x annual fees |" in table
    assert "| --- | --- |" in table

    prose = layout_to_markdown("The cap is 2x annual fees.\nSee section 9.3.\n")
    assert "| " not in prose
    assert "2x annual fees" in prose

    spaced = layout_to_markdown(
        "The cap is 2x.  See section 9.3.\nThe window is 14 days.  No exceptions.\n"
    )
    assert "| " not in spaced
